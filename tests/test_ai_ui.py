"""Real Streamlit execution with a mocked API; no paid requests or real account records."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from ai_client import AIConfig,AIError,request_version,MISSING
from ai_state import new_workspace,install_result,accept,undo_accept,export_ai,reopen_chat,append_chat
from ai_report import report_tables,markdown_appendix
from ai_summary import build_summary
from test_ai_core import fixture,response_for
from project_store import save_project,load_project
from project_state import new_project
from dynamic_topics import discover_topics
import excel_export

ROOT=Path(__file__).parents[1]


def fake_response(summary,config,question=None,history=None):
    return {'result':response_for(summary,question is not None),'summary':deepcopy(summary),
            'version':request_version(summary,config),'model':config.model,'returned_model':config.model,
            'generated_at':'2026-09-30T00:00:00+00:00','prompt_version':'test','contract_version':'test','app_version':'V3','usage':None}


def widget(app,kind,label): return next(item for item in getattr(app,kind) if item.label==label)


def loaded_app():
    app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=45).run()
    app.button('cover-start').click().run()
    snapshot,_=fixture(); snapshot['url']='https://www.xiaohongshu.com/user/profile/testmbti'
    snapshot.update(profile_text='模拟测试简介',stop_reason='测试样本已加载')
    app.session_state['public_snapshot']={'input':snapshot['url'],'data':snapshot}
    app.button('source-card-2').click().run()
    return app,snapshot


class AIStateTests(unittest.TestCase):
    def test_adoption_drafts_staleness_chat_and_project_roundtrip(self):
        snap,analysis=fixture(); summary=build_summary(snap,analysis); config=AIConfig()
        record=fake_response(summary,config); version=record['version']; work=new_workspace(summary['account_id'])
        install_result(work,record)
        self.assertEqual(export_ai(work,version),{})
        self.assertIn('draft',export_ai(work,version,True))
        edited=deepcopy(work['edited']); edited['profile']='人工采纳的混合定位'
        accept(work,version,edited)
        self.assertEqual(analysis['position'],'秘密内部备注')
        bundle=export_ai(work,version)
        self.assertIn('人工采纳的混合定位',markdown_appendix(bundle))
        self.assertIn('人工采纳的混合定位',report_tables(bundle)['AI采纳结论']['结论'].tolist())
        self.assertEqual(export_ai(work,'new-version',True),{})
        undo_accept(work); self.assertEqual(export_ai(work,version),{})
        undo_accept(work); self.assertIn('accepted',export_ai(work,version))
        reopen_chat(work,version); append_chat(work,version,'测试问题',fake_response(summary,config,'测试问题'))
        with self.assertRaises(ValueError): append_chat(work,'new-version','问题',record)
        reopen_chat(work,'new-version'); self.assertEqual(work['chat'],[]); self.assertEqual(len(work['archived_chat']),1)
        with TemporaryDirectory() as temp:
            project=new_project(snap,discover_topics(snap['notes'])); project['state']['ai_workspace']=work
            saved=save_project(project,Path(temp)); restored=load_project(saved['id'],Path(temp))
            self.assertEqual(restored['state']['ai_workspace']['accepted']['sections']['profile'],'人工采纳的混合定位')


class AIUITests(unittest.TestCase):
    def test_no_key_and_no_implicit_request(self):
        with patch('ai_ui.load_config',return_value=AIConfig()),patch('ai_ui.ask_ai') as api:
            app,_=loaded_app()
            self.assertFalse(app.exception)
            self.assertTrue(any(i.value==MISSING for i in app.info))
            self.assertTrue(widget(app,'button','生成AI分析').disabled)
            widget(app,'text_area','你的账号定位结论').set_value('保留人工定位').run()
            app.run()
            api.assert_not_called()
            self.assertFalse(app.exception)

    def test_generate_edit_adopt_failure_chat_stale_and_export(self):
        captured={}; original=excel_export.public_workbook
        def export(snapshot,analysis):
            data=original(snapshot,analysis); captured.update(analysis=deepcopy(analysis),data=data); return data
        with TemporaryDirectory() as temp,patch('project_ui.STORE_ROOT',Path(temp)),patch('ai_ui.load_config',return_value=AIConfig(api_key='fake')),patch('ai_ui.ask_ai',side_effect=fake_response) as api,patch('excel_export.public_workbook',side_effect=export):
            app,snapshot=loaded_app()
            widget(app,'text_area','你的账号定位结论').set_value('已有人工定位').run()
            api.assert_not_called()
            widget(app,'button','生成AI分析').click().run()
            self.assertFalse(app.exception,app.exception)
            self.assertEqual(api.call_count,1)
            app.run(); self.assertEqual(api.call_count,1)
            self.assertEqual(captured['analysis']['ai'],{})
            widget(app,'text_area','账号识别与定位').set_value('人工编辑后采纳的结论')
            widget(app,'button','保存编辑并采纳').click().run()
            self.assertFalse(app.exception,app.exception)
            self.assertEqual(widget(app,'text_area','你的账号定位结论').value,'已有人工定位')
            adopted=captured['analysis']['ai']['accepted']
            self.assertEqual(adopted['sections']['profile'],'人工编辑后采纳的结论')
            book=load_workbook(BytesIO(captured['data']))
            self.assertIn('AI采纳结论',book.sheetnames)
            self.assertNotIn('AI草稿附录结论',book.sheetnames)
            self.assertIn('人工编辑后采纳的结论',[cell.value for row in book['AI采纳结论'] for cell in row])
            # A failed regenerate retains draft and adoption.
            api.side_effect=AIError('timeout','模拟超时')
            widget(app,'button','重新生成AI分析').click().run()
            self.assertEqual(captured['analysis']['ai']['accepted'],adopted)
            self.assertTrue(any('模拟超时' in e.value for e in app.error))
            api.side_effect=fake_response
            widget(app,'button','重开当前数据对话').click().run()
            widget(app,'text_input','向AI提问').set_value('最近是否下降？')
            widget(app,'button','发送问题').click().run()
            self.assertFalse(app.exception,app.exception)
            self.assertEqual(api.call_count,3)
            widget(app,'button','保存项目').click().run()
            self.assertFalse(app.exception,app.exception)
            project=next(Path(temp).glob('*.json'))
            saved=load_project(project.stem,Path(temp))
            self.assertEqual(saved['state']['ai_workspace']['accepted']['sections']['profile'],'人工编辑后采纳的结论')
            # Fresh session restores AI state without a request.
            restored=AppTest.from_file(str(ROOT/'app.py'),default_timeout=45).run()
            restored.button('cover-start').click().run()
            restored.session_state['active_project']=saved
            restored.session_state['show_saved_project']=True
            restored.session_state['restore_epoch']=1
            restored.run()
            self.assertFalse(restored.exception,restored.exception)
            self.assertIn('accepted',captured['analysis']['ai'])
            self.assertEqual(api.call_count,3)
            widget(app,'multiselect','运营目标（由你确认，可多选）').set_value(['咨询']).run()
            self.assertEqual(captured['analysis']['ai'],{})
            self.assertTrue(widget(app,'button','发送问题').disabled)
            self.assertTrue(any('已过期' in w.value for w in app.warning))
            self.assertEqual(api.call_count,3)
            other=deepcopy(snapshot); other.update(name='另一个账号',account_id='another',url='https://www.xiaohongshu.com/user/profile/another')
            app.session_state['public_snapshot']={'input':other['url'],'data':other}; app.run()
            self.assertFalse(app.exception,app.exception)
            self.assertEqual(captured['analysis']['ai'],{})
            self.assertTrue(any(b.label=='生成AI分析' for b in app.button))
            self.assertEqual(api.call_count,3)


if __name__=='__main__': unittest.main()
