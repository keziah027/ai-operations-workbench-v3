from copy import deepcopy
from io import BytesIO
import unittest

import pandas as pd
from openpyxl import load_workbook

from content_analysis import build_annotations, build_detail, pillar_summary, audience_signals, high_interaction
from excel_export import public_workbook, creator_workbook, Report


def fixture():
    notes = [{"标题": f"新品教程{i}", "点赞": i * 10, "收藏": i, "评论": i % 3,
              "发布时间": "2026-09-24 12:00", "正文": "准备贴纸\n1. 拼贴\n2. 到店铺购买",
              "笔记链接": f"https://www.xiaohongshu.com/explore/{i:024x}"} for i in range(12)]
    notes[0]["标题"] = '=HYPERLINK("https://example.com")'
    notes[1]["收藏"] = None
    notes[1]["发布时间"] = None
    snapshot = {"name": "测试账号", "captured_at": "2026-09-25T12:00:00+08:00", "notes": notes,
                "bio": "原创手账", "url": "https://www.xiaohongshu.com/user/profile/aaaaaaaaaaaaaaaaaaaaaaaa",
                "stop_reason": "本次测试", "requested_limit": 100}
    annotations = build_annotations(notes)
    annotations.loc[0, "主支柱"] = "用户共创"
    annotations.loc[0, "封面观察"] = "人工封面观察"
    annotations.loc[0, "商品/系列（人工确认）"] = "春日系列"
    detail = build_detail(notes, annotations)
    valid, top, comparisons = high_interaction(detail, "收藏")
    analysis = {"position": "人工定位结论", "audience": "人工人群假设", "audience_findings": audience_signals(detail, snapshot["bio"]),
                "metric": "收藏", "valid": valid, "top": top, "comparisons": comparisons,
                "summary": pillar_summary(detail), "detail": detail, "annotations": annotations,
                "commerce": [{"承接线索": "店铺", "命中篇数": 12, "占全部样本（%）": 100}]}
    return snapshot, analysis


def rows(sheet):
    headers = [cell.value for cell in sheet[4]]
    return [{key: cell for key, cell in zip(headers, row)} for row in sheet.iter_rows(min_row=5)]


class ExcelExportTests(unittest.TestCase):
    def test_public_snapshot_manual_edits_missing_and_zero(self):
        snapshot, analysis = fixture()
        before = deepcopy(snapshot)
        book = load_workbook(BytesIO(public_workbook(snapshot, analysis)))
        self.assertEqual(snapshot, before)
        self.assertEqual(len(book.sheetnames), 5)
        notes = rows(book["笔记明细"])
        self.assertEqual(len(notes), 12)
        self.assertEqual(notes[0]["点赞"].value, 0)
        self.assertIsNone(notes[0]["藏赞比"].value)
        self.assertIsNone(notes[1]["收藏"].value)
        self.assertIsNone(notes[1]["发布时间"].value)
        self.assertIsNone(notes[1]["互动总量"].value)
        self.assertAlmostEqual(notes[2]["藏赞比"].value, .1)
        self.assertEqual(notes[2]["藏赞比"].number_format, "0.00%")
        self.assertEqual(notes[0]["标题"].data_type, "s")
        detail = rows(book["内容拆解明细"])
        self.assertEqual(detail[0]["主支柱"].value, "用户共创")
        self.assertEqual(detail[0]["封面观察"].value, "人工封面观察")
        self.assertEqual(detail[0]["商品/系列（人工确认）"].value, "春日系列")
        values = [cell.value for row in book["五维分析结果"] for cell in row]
        self.assertIn("人工定位结论", values)
        self.assertIn("人工人群假设", values)
        self.assertEqual(len(book["五维分析结果"]._charts), 1)
        self.assertTrue(all(cell.data_type != "f" for sheet in book for row in sheet for cell in row))

    def test_creator_overview_and_ranking(self):
        snapshot, _ = fixture()
        snapshot.update(overview={"name": "我的账号", "period": "近7日", "metrics": [{"指标": "曝光", "后台显示": "1.4万"}]},
                        filters={"start": "2026-09-01", "end": "2026-09-25", "type": "全部"}, complete_filtered=False)
        book = load_workbook(BytesIO(creator_workbook(snapshot)))
        self.assertEqual(book["账号期间总览"]["B5"].value, "1.4万")
        self.assertEqual(len(rows(book["排行榜"])), 30)
        self.assertEqual(rows(book["排行榜"])[0]["指标数值"].value, 110)
        self.assertEqual(len(book["发布时间趋势"]._charts), 1)

    def test_partial_overview_without_notes(self):
        book = load_workbook(BytesIO(creator_workbook({"overview": {"metrics": [{"指标": "曝光", "后台显示": "1.4万"}]}, "notes": []})))
        self.assertIn("账号期间总览", book.sheetnames)
        self.assertIn("没有可导出的记录", book["笔记明细"]["A5"].value)

    def test_long_text_and_nat(self):
        report = Report()
        report.table(report.sheet("原始数据"), "数据", pd.DataFrame({"正文": ["长" * 33000], "日期": [pd.NaT]}))
        book = load_workbook(BytesIO(report.finish()))
        self.assertIsNone(book["原始数据"]["B5"].value)
        self.assertEqual(book["长文本补充"]["D5"].value + book["长文本补充"]["D6"].value, "长" * 33000)


if __name__ == "__main__":
    unittest.main()
