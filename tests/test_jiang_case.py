"""The replacement must remain unpinned through analysis and exports."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import unittest
import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from real_cases import list_cases, validate_case
from note_schema import normalize_notes
from glass_ui import overview_metric, trend_data


class JiangCaseTests(unittest.TestCase):
    def case(self):
        return next(c for c in list_cases() if c['name']=='姜乘澜')

    def test_pinned_and_unverified_notes_cannot_enter_analysis(self):
        case=self.case()
        self.assertEqual(len(case['notes']),18)
        self.assertFalse(any(c['name'].startswith('程十安') for c in list_cases()))
        for mutation in ('marked','missing','excluded'):
            invalid=deepcopy(case)
            if mutation=='marked': invalid['notes'][0]['是否置顶']=True
            elif mutation=='missing': invalid['notes'][0].pop('置顶核验时间')
            else: invalid['notes'][0]['笔记链接']=invalid['excluded_pinned'][0]['url']
            with self.assertRaises(ValueError): validate_case(invalid)

    def test_public_like_display_does_not_invent_exact_likes(self):
        frame=pd.DataFrame(normalize_notes(self.case()['notes']))
        self.assertEqual(overview_metric(frame,'评论'),'点赞')
        self.assertTrue(frame['点赞'].isna().all())
        self.assertEqual(frame['评论'].count(),17)
        self.assertEqual(frame['评论'].sum(),80519)
        self.assertFalse(trend_data(frame,'评论').empty)
        existing=pd.DataFrame(normalize_notes(list_cases()[0]['notes']))
        self.assertEqual(overview_metric(existing,'评论'),'点赞')

    def test_ui_and_exports_use_only_the_replacement_sample(self):
        import excel_export
        import content_ui_v2
        original_excel=excel_export.public_workbook
        original_markdown=content_ui_v2.markdown_report
        captured={}
        def excel(snapshot,analysis):
            data=original_excel(snapshot,analysis)
            captured.update(excel=data,snapshot=snapshot)
            return data
        def markdown(snapshot,analysis):
            data=original_markdown(snapshot,analysis)
            captured['markdown']=data
            return data
        with patch('excel_export.public_workbook',side_effect=excel), patch('content_ui_v2.markdown_report',side_effect=markdown):
            app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=90).run()
            app.button('cover-start').click().run()
            app.button('source-card-1').click().run()
            app.selectbox('source-case-choice').select(self.case()['id']).run()
            next(b for b in app.button if b.label=='加载真实账号案例').click().run()
        self.assertFalse(app.exception,app.exception)
        self.assertEqual(app.session_state['overview-metric'],'点赞')
        self.assertEqual(captured['snapshot']['name'],'姜乘澜')
        self.assertEqual(len(captured['snapshot']['notes']),18)
        workbook=load_workbook(BytesIO(captured['excel']),read_only=True,data_only=True)
        cells=[str(v) for sheet in workbook for row in sheet.iter_rows(values_only=True) for v in row if v is not None]
        for note in self.case()['notes']:
            self.assertIn(note['标题'],cells)
        # Source-limit prose may name an excluded work; note rows must not contain it.
        self.assertNotIn('好久不见',cells)
        self.assertFalse(any('程十安' in c for c in cells))
        self.assertIn('姜乘澜',captured['markdown'])
        self.assertIn('点赞：10万+',captured['markdown'])
        self.assertIn('10万+',cells)
        self.assertIn('4.1万',cells)
        self.assertNotIn('程十安',captured['markdown'])
        self.assertNotIn(self.case()['excluded_pinned'][0]['url'],captured['markdown'])
        app.button('source-back').click().run()
        app.button('source-card-1').click().run()
        self.assertEqual(app.session_state['active-real-case-id'],self.case()['id'])
        self.assertFalse(app.exception)


if __name__=='__main__': unittest.main()
