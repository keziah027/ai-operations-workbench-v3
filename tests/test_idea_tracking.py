import unittest
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from analysis import prepare_data, idea_effects, read_excel
from idea_tracking import adopted_idea_pool
from excel_export import public_workbook
from test_excel_export import fixture
from test_ai_core import fixture as ai_fixture
from test_ai_ui import fake_response
from ai_client import AIConfig
from ai_summary import build_summary
from ai_state import new_workspace, install_result, accept, export_ai, undo_accept


def accepted_bundle():
    snap,analysis=ai_fixture();summary=build_summary(snap,analysis)
    record=fake_response(summary,AIConfig());workspace=new_workspace(summary['account_id'])
    install_result(workspace,record)
    workspace['edited']['ideas']='【1】\n\n标题草稿：人工修改后的标题\n\n建议形式：图文\n\n所属主题或新方向：职业发展\n\n核心思路：第一行\n第二行\n\n观察指标：收藏与评论\n\n证据：保留原文'
    accept(workspace,record['version'],workspace['edited'])
    return workspace,record['version']


class IdeaTrackingTests(unittest.TestCase):
    def test_labels_alias_placeholders_and_no_implicit_linking(self):
        df,_=prepare_data(pd.DataFrame({'标题':list('abcdef'),'选题编号':['XT-abc',None,'—','未分类',' nan ','  自定义方案  ']}))
        self.assertEqual(df['对应选题'].tolist(),['XT-abc','未关联','未关联','未关联','未关联','自定义方案'])
        self.assertEqual(idea_effects(df)['对应选题'].tolist(),['XT-abc','自定义方案'])
        self.assertTrue(idea_effects(pd.DataFrame()).empty)
        self.assertTrue(idea_effects(pd.DataFrame({'对应选题':[None,'未关联']})).empty)

    def test_same_account_threshold_and_unlinked_excluded(self):
        df,_=prepare_data(pd.DataFrame({'标题':[str(i) for i in range(12)], '账号名称':['甲']*12,
            '对应选题':['低组']*4+['高组']*4+[None]*4,
            '点赞':[1,2,3,4,90,100,110,'1.2万',5,6,7,8], '收藏':[None]*12}))
        result=idea_effects(df).set_index('对应选题')
        self.assertEqual(result.index.tolist(),['低组','高组'])
        self.assertEqual(result.loc['低组','爆款率'],0)
        self.assertTrue(pd.isna(result.loc['高组','爆款率']))
        self.assertEqual(result.loc['高组','精确点赞篇数'],3)
        self.assertEqual(result['账号有效篇数'].tolist(),[11,11])
        self.assertTrue(result['平均收藏'].isna().all())

    def test_multi_account_no_shared_baseline(self):
        a=pd.DataFrame({'对应选题':['同一编号']*4,'账号名称':['甲']*4,'点赞':[1,2,3,4]})
        b=pd.DataFrame({'对应选题':['同一编号']*4,'账号名称':['乙']*4,'点赞':[100,200,300,400]})
        r=idea_effects(pd.concat([a,b],ignore_index=True))
        self.assertEqual(r['账号上四分位'].tolist(),[3.25,325])
        self.assertEqual(r['爆款率'].tolist(),[.25,.25])

    def test_adopted_edited_text_stable_ids_and_stale_draft_excluded(self):
        work,version=accepted_bundle();bundle=export_ai(work,version)
        result=adopted_idea_pool(bundle)
        self.assertEqual(result.iloc[0]['标题草稿'],'人工修改后的标题')
        self.assertEqual(result.iloc[0]['核心思路'],'第一行\n第二行')
        self.assertEqual(result.iloc[0]['观察指标'],'收藏与评论')
        pd.testing.assert_frame_equal(result,adopted_idea_pool(deepcopy(bundle)))
        self.assertTrue(adopted_idea_pool(export_ai(work,'旧版本',True)).empty)
        undo_accept(work)
        self.assertTrue(adopted_idea_pool(export_ai(work,version,True)).empty)
        self.assertTrue(adopted_idea_pool({}).empty)

    def test_freely_edited_text_preserved_without_invented_fields(self):
        text='不按卡片格式写的人工选题\n第二行'
        r=adopted_idea_pool({'accepted':{'sections':{'ideas':text}}})
        self.assertEqual(r.iloc[0]['核心思路'],text)
        self.assertEqual(r.iloc[0]['采纳选题原文'],text)
        self.assertEqual(r.iloc[0]['建议形式'],'')

    def test_export_pool_only_adopted_formula_safe_and_link_roundtrip(self):
        snap,analysis=fixture();work,version=accepted_bundle()
        bundle=export_ai(work,version);bundle['accepted']['sections']['ideas']=bundle['accepted']['sections']['ideas'].replace('人工修改后的标题','=1+1')
        analysis['ai']=bundle
        with patch('ai_client.ask_ai',side_effect=AssertionError('Export must not call AI')):
            book=load_workbook(BytesIO(public_workbook(snap,analysis)))
        self.assertIn('选题池',book.sheetnames)
        headers=[cell.value for cell in book['选题池'][4]]
        row=dict(zip(headers,book['选题池'][5]))
        self.assertEqual(row['标题草稿'].value,'=1+1')
        self.assertEqual(row['标题草稿'].data_type,'s')
        identity=row['选题编号'].value
        source=BytesIO();pd.DataFrame({'标题':list('abcd'),'选题编号':[identity]*4,'点赞':[10,20,30,40]}).to_excel(source,index=False)
        cleaned,_=read_excel(source.getvalue())
        effects=idea_effects(cleaned)
        self.assertEqual(effects.iloc[0]['对应选题'],identity)
        self.assertEqual(effects.iloc[0]['爆款率'],.25)
        analysis['idea_effects']=effects
        self.assertIn('选题效果',load_workbook(BytesIO(public_workbook(snap,analysis))).sheetnames)
        analysis['ai']={'draft':bundle['accepted']}
        self.assertNotIn('选题池',load_workbook(BytesIO(public_workbook(snap,analysis))).sheetnames)
