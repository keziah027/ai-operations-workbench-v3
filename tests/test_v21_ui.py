"""真实页面执行路径与下载内容的一致性（无登录、无网络读取）。"""
from pathlib import Path
from io import BytesIO
import hashlib
import json
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
import excel_export
from cross_account_cases import CASES, case_notes
from note_schema import normalize_notes
from dynamic_topics import discover_topics, OTHER


class UIRegressionTests(unittest.TestCase):
    def test_axes_manual_edits_sampling_backend_and_exports(self):
        captured={}
        original=excel_export.public_workbook
        def export(snapshot,analysis):
            data=original(snapshot,analysis)
            captured.update(snapshot=snapshot,analysis=analysis,data=data)
            return data
        def widget(app,kind,label):
            return next(item for item in getattr(app,kind) if item.label==label)
        with patch('excel_export.public_workbook',side_effect=export):
            app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=45).run()
            app.button('cover-start').click().run()
            for index,case in enumerate(CASES):
                notes=case_notes(case)
                notes[0]['点赞']='1.2万'
                snapshot={'name':case['name'],'url':'https://www.xiaohongshu.com/user/profile/aaaaaaaaaaaaaaaaaaaaaaaa',
                          'bio':case['bio'],'captured_at':f'2026-09-26T12:00:0{index}+08:00','profile_text':'测试数据',
                          'notes':notes,'stop_reason':'测试完成'}
                app.session_state['public_snapshot']={'input':snapshot['url'],'data':snapshot}
                if index:
                    app.button('source-back').click().run()
                app.button('source-card-2').click().run()
                self.assertFalse(app.exception,app.exception)
                self.assertEqual(captured['analysis']['goals'],[])
                widget(app,'multiselect','内容领域（可多选）').set_value(case['domains']).run()
                widget(app,'multiselect','运营目标（由你确认，可多选）').set_value(case['goals']).run()
                widget(app,'text_area','你的账号定位结论').set_value('人工核对：'+case['name']).run()
                self.assertFalse(app.exception,app.exception)
                analysis=captured['analysis']
                self.assertEqual(analysis['domains'],case['domains'])
                self.assertEqual(analysis['goals'],case['goals'])
                self.assertEqual(analysis['sample_audit']['是否纳入'].eq('否').sum(),1)
                book=load_workbook(BytesIO(captured['data']))
                values=[cell.value for row in book['五模块分析结果'] for cell in row]
                self.assertIn('人工核对：'+case['name'],values)
                self.assertIn('、'.join(case['goals']),values)
                self.assertIn('、'.join(case['domains']),values)
                self.assertIn('比较样本范围',book.sheetnames)
                self.assertIn('后台支柱表现',book.sheetnames)
                headers=[cell.value for cell in book['笔记明细'][4]]
                self.assertNotIn('收藏率',headers)
                self.assertIn('藏赞比',headers)
                self.assertEqual(book['笔记明细'].cell(5,headers.index('点赞原文')+1).value,'1.2万')
                self.assertEqual(book['笔记明细'].cell(5,headers.index('点赞')+1).value,'1.2万')
                self.assertIsNone(book['笔记明细'].cell(5,headers.index('点赞（精确值）')+1).value)
            # Edit a note and then change the commerce gate; manual values survive the editor reconfiguration.
            normal=normalize_notes(notes)
            key=hashlib.sha256(json.dumps([snapshot['url'],snapshot['captured_at'],normal],ensure_ascii=False,default=str).encode()).hexdigest()[:16]
            labels=captured['analysis']['topics']['内容支柱'].tolist()
            editor=f'annotations-v3-{key}-0-{int(captured['analysis']['commerce_enabled'])}'
            edits={'edited_rows':{0:{'主支柱':labels[1],'封面观察':'保留人工标注'}},'added_rows':[],'deleted_rows':[]}
            app.session_state[editor]=edits
            app.run()
            self.assertEqual(captured['analysis']['detail'].iloc[0]['封面观察'],'保留人工标注')
            # AppTest does not simulate a data-editor client; resend its unchanged delta.
            widget(app,'multiselect','运营目标（由你确认，可多选）').set_value(['到店','销售'])
            app.session_state[editor]=edits
            app.run()
            self.assertFalse(app.exception,app.exception)
            self.assertEqual(captured['analysis']['detail'].iloc[0]['主支柱'],labels[1])
            self.assertEqual(captured['analysis']['detail'].iloc[0]['封面观察'],'保留人工标注')
            revised_book=load_workbook(BytesIO(captured['data']))
            revised_sheet=revised_book['内容拆解明细']
            revised_headers=[cell.value for cell in revised_sheet[4]]
            self.assertEqual(revised_sheet.cell(5,revised_headers.index('主支柱')+1).value,labels[1])
            self.assertEqual(revised_sheet.cell(5,revised_headers.index('封面观察')+1).value,'保留人工标注')
            widget(app,'selectbox','用哪项数据比较高互动内容').set_value('每千次观看涨粉').run()
            self.assertFalse(app.exception,app.exception)
            self.assertEqual(captured['analysis']['sample_audit']['是否纳入'].eq('是').sum(),8)
            # The creator path keeps the distinct account-period overview and the same core fields.
            creator={'overview':{'name':'模拟自有账号','account_id':'test','period':'近7日',
                        'metrics':[{'指标':'曝光数','后台显示':'1.4万','环比':'环比-7%'}]},
                     'notes':case_notes(CASES[2]),'captured_at':'2026-09-26T13:00:00+08:00','days':7,
                     'filters':{'start':'2026-09-01','end':'2026-09-26','type':'全部'},
                     'platform_total':8,'complete_filtered':True,'stop_reason':'测试完成','warnings':[],
                     'source':'https://creator.xiaohongshu.com/statistics/data-analysis'}
            app.session_state['creator_snapshot']=creator
            app.button('source-back').click().run()
            app.button('source-card-3').click().run()
            self.assertFalse(app.exception,app.exception)
            self.assertEqual(captured['analysis']['goals'],[])
            self.assertEqual(captured['analysis']['position'],'')
            self.assertEqual(captured['snapshot']['source_kind'],'自有账号后台')
            detail=captured['analysis']['detail']
            self.assertEqual(detail['曝光'].count(),8)
            self.assertEqual(detail['内容形式'].ne('').sum(),8)
            book=load_workbook(BytesIO(captured['data']))
            self.assertEqual(book['账号期间总览']['B5'].value,'1.4万')
            self.assertEqual(book['笔记明细'].max_row,12)
            self.assertIn('发布时间表现',book.sheetnames)


if __name__=='__main__': unittest.main()
