"""Exercise the real app's selection routes, availability, and return state."""
import os
from pathlib import Path
import unittest
from unittest.mock import patch
from streamlit.testing.v1 import AppTest

APP = Path(__file__).parents[1] / 'app.py'


class SourcePickerTests(unittest.TestCase):
    def app(self):
        app = AppTest.from_file(str(APP), default_timeout=60).run()
        app.button('cover-start').click().run()
        self.assertEqual(len(app.get('popover')), 0)
        self.assertEqual(len(app.radio), 0)
        return app

    def clean(self, app):
        self.assertFalse(app.exception, app.exception)

    @patch.dict(os.environ, {'APP_MODE': 'public'})
    def test_public_routes_keep_real_availability_messages(self):
        app = self.app()
        self.clean(app)
        self.assertEqual([b.label for b in app.button if b.key and b.key.startswith('source-card-')],
                         ['Excel 文件', '真实账号案例', '公开账号', '我的账号后台'])
        for index, expected in [(2, '主页读取仅支持本地版'), (3, '后台同步仅支持本地版')]:
            app.button(f'source-card-{index}').click().run()
            self.clean(app)
            self.assertTrue(any(expected in message.value for message in app.markdown))
            self.assertFalse(app.session_state['source-picker-open'])
            app.button('source-back').click().run()
            self.assertTrue(app.session_state['source-picker-open'])
        app.button('source-card-0').click().run()
        self.clean(app)
        self.assertTrue(any(item.label == '选择演示账号类型' for item in app.selectbox))

    @patch.dict(os.environ, {'APP_MODE': 'public'})
    @patch('real_cases.list_cases', return_value=[])
    def test_missing_cases_show_truthful_empty_state(self, _cases):
        app = self.app()
        app.button('source-card-1').click().run()
        self.clean(app)
        self.assertTrue(any('暂无可用的真实案例快照' in message.value for message in app.info))
        self.assertFalse(any(b.label == '加载真实账号案例' for b in app.button))

    @patch.dict(os.environ, {'APP_MODE': 'public'})
    def test_non_default_case_and_analysis_session_survive_return(self):
        app = self.app()
        app.button('source-card-1').click().run()
        choice = app.selectbox('source-case-choice')
        from real_cases import list_cases
        selected = list_cases()[1]['id']
        choice.select(selected).run()
        case_id = app.session_state['source-case-choice']
        app.session_state['active-real-case-id'] = case_id
        app.session_state['editor-state-regression-sentinel'] = {'annotations': ['retained']}
        app.button('source-back').click().run()
        app.button('source-card-1').click().run()
        self.clean(app)
        self.assertEqual(app.session_state['source-case-choice'], case_id)
        self.assertEqual(app.session_state['active-real-case-id'], case_id)
        self.assertEqual(app.session_state['editor-state-regression-sentinel'], {'annotations': ['retained']})


    @patch.dict(os.environ, {'APP_MODE': 'public'})
    def test_all_cases_load_and_export_without_route_changes(self):
        import excel_export
        from real_cases import list_cases
        from io import BytesIO
        from openpyxl import load_workbook
        captured = {}
        original = excel_export.public_workbook
        def export(snapshot, analysis):
            result = original(snapshot, analysis)
            captured.update(name=snapshot['name'], notes=snapshot['notes'], excel=result)
            return result
        with patch('excel_export.public_workbook', side_effect=export):
            app = self.app()
            # A browser may retain the old source radio's Excel value.
            app.session_state['data-source'] = 'Excel 文件'
            app.button('source-card-1').click().run()
            for index, case in enumerate(list_cases()):
                if index:
                    app.selectbox('source-case-choice').select(case['id']).run()
                next(b for b in app.button if b.label == '加载真实账号案例').click().run()
                self.clean(app)
                self.assertEqual(app.session_state['active-data-source'], '真实账号案例')
                self.assertEqual(captured['name'], case['name'])
                self.assertEqual(len(captured['notes']), 18)
                self.assertFalse(any(s.label == '选择演示账号类型' for s in app.selectbox))
                app.button('open-analysis-report').click().run()
                self.assertTrue(app.session_state['analysis-report-open'])
                panel = next(e for e in app.expander if e.label == '内容分析与报告')
                self.assertTrue(panel.proto.expanded)
                next(b for b in app.button if b.label == '保存项目').click().run()
                self.clean(app)
                self.assertFalse(app.session_state['project_dirty'])
                self.assertEqual(app.session_state['active-data-source'], '真实账号案例')
                book = load_workbook(BytesIO(captured['excel']), read_only=True)
                self.assertEqual(book['笔记明细'].max_row, 22)
                order = next(s for s in app.selectbox if s.label == '时间排序')
                order.select(order.options[-1]).run()
                app.run()
                self.assertEqual(app.session_state['active-data-source'], '真实账号案例')
                app.button('source-back').click().run()
                self.assertEqual(len(app.get('popover')), 0)
                app.button('source-card-1').click().run()
                self.assertEqual(app.selectbox('source-case-choice').value, case['id'])
                self.assertEqual(captured['name'], case['name'])

    @patch.dict(os.environ, {'APP_MODE': 'public'})
    def test_demo_and_case_remain_separate_across_navigation(self):
        app = self.app()
        app.button('source-card-0').click().run()
        next(b for b in app.button if b.label == '加载演示数据').click().run()
        scenario = app.session_state['demo_scenario']
        for index in (1, 2, 3):
            app.button('source-back').click().run()
            app.button(f'source-card-{index}').click().run()
            self.clean(app)
            self.assertFalse(any(s.label == '选择演示账号类型' for s in app.selectbox))
            self.assertEqual(app.session_state['demo_scenario'], scenario)
        app.button('source-back').click().run()
        app.button('source-card-0').click().run()
        self.clean(app)
        self.assertEqual(app.session_state['active-data-source'], 'Excel 文件')
        self.assertEqual(next(m.value for m in app.metric if m.label == '总发布篇数'), '18')


if __name__ == '__main__':
    unittest.main()
