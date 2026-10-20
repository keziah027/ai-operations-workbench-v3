"""Public deployment acceptance tests, run separately from desktop regression tests."""
from io import BytesIO
from pathlib import Path
from unittest.mock import patch, Mock
import json
import os
import unittest
import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from analysis import read_excel, summarize
from demo_data import SCENARIOS, demo_workbook
from ai_client import AIConfig, ask_ai, load_config

APP = str(Path(__file__).resolve().parents[1] / 'app.py')


def widget(app, kind, label):
    return next(item for item in getattr(app, kind) if item.label == label)


class PublicTests(unittest.TestCase):
    def setUp(self):
        self.assertNotEqual(os.environ.get('APP_MODE'), 'local')

    def test_four_scenarios_cleaning_and_stable_workbooks(self):
        for scenario in SCENARIOS:
            with self.subTest(scenario=scenario):
                contents = demo_workbook(scenario)
                self.assertEqual(contents, demo_workbook(scenario))
                frame, messages = read_excel(contents)
                self.assertEqual(len(frame), 18)
                self.assertEqual(frame.attrs['source']['header_row'], 3)
                self.assertTrue(pd.isna(frame.loc[2, '点赞']))
                self.assertEqual(frame.loc[2, '点赞精度'], '近似')
                self.assertTrue(pd.isna(frame.loc[5, '收藏']))
                self.assertEqual(frame.loc[8, '评论精度'], '下界')
                self.assertTrue(pd.isna(frame.loc[11, '藏赞比']))
                self.assertTrue(frame['数据性质'].str.contains('模拟数据').all())
                self.assertEqual(summarize(frame)['总发布篇数'], 18)
                self.assertGreaterEqual(len(messages), 3)

    def test_upload_limits_and_invalid_data(self):
        with self.assertRaises(ValueError):
            read_excel(b'not an excel file')
        with self.assertRaises(ValueError):
            read_excel(b'x' * (10 * 1024 * 1024 + 1))
        output = BytesIO()
        pd.DataFrame({'标题': ['模拟数据'] * 1001}).to_excel(output, index=False)
        with self.assertRaisesRegex(ValueError, '1000'):
            read_excel(output.getvalue())

    def test_end_to_end_all_scenarios_and_exports(self):
        from excel_export import public_workbook
        captured = {}
        def capture(snapshot, analysis):
            report = public_workbook(snapshot, analysis)
            captured.update(snapshot=snapshot, analysis=analysis, report=report)
            return report
        with patch('excel_export.public_workbook', side_effect=capture):
            for scenario in SCENARIOS:
                app = AppTest.from_file(APP, default_timeout=90).run()
                app.button('cover-start').click().run()
                app.button('source-card-0').click().run()
                self.assertFalse(app.exception, app.exception)
                widget(app, 'selectbox', '选择演示账号类型').set_value(scenario).run()
                widget(app, 'button', '加载演示数据').click().run()
                self.assertFalse(app.exception, app.exception)
                self.assertTrue(captured['snapshot']['is_demo'])
                self.assertEqual(len(captured['analysis']['detail']), 18)
                self.assertIn('模拟', captured['snapshot']['source_kind'])
                book = load_workbook(BytesIO(captured['report']))
                self.assertGreater(len(book.sheetnames), 10)
                all_text = '\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)
                self.assertIn('模拟', all_text)
                widget(app, 'button', '生成AI分析').click().run()
                self.assertFalse(app.exception, app.exception)
                widget(app, 'button', '保存编辑并采纳').click().run()
                self.assertFalse(app.exception, app.exception)
                self.assertIn('accepted', captured['analysis']['ai'])
                self.assertIn('规则演示', captured['analysis']['ai']['accepted']['model'])
                widget(app, 'text_input', '向AI提问').set_value('怎样改善内容表现？')
                widget(app, 'button', '发送问题').click().run()
                self.assertFalse(app.exception, app.exception)
                widget(app, 'button', '保存项目').click().run()
                self.assertFalse(app.exception, app.exception)
                project = app.session_state['active_project']
                from project_store import validate
                validate(project)
                self.assertIn('ai_workspace', project['state'])

    def test_sessions_isolated_and_cloud_browser_unavailable(self):
        app = AppTest.from_file(APP, default_timeout=90).run()
        app.button('cover-start').click().run()
        app.button('source-card-0').click().run()
        widget(app, 'button', '加载演示数据').click().run()
        widget(app, 'button', '保存项目').click().run()
        self.assertEqual(len(app.session_state['_private_projects']), 1)
        fresh = AppTest.from_file(APP, default_timeout=90).run()
        fresh.button('cover-start').click().run()
        fresh.button('source-card-0').click().run()
        self.assertEqual(len(fresh.session_state['_private_projects']), 0)
        self.assertEqual(widget(fresh, 'selectbox', '选择已保存项目').options, ['请选择项目'])
        for index in (2, 3):
            fresh.button('source-back').click().run()
            fresh.button(f'source-card-{index}').click().run()
            self.assertFalse(fresh.exception, fresh.exception)
            self.assertFalse(any('登录窗口' in button.label for button in fresh.button))

    def test_paid_api_blocked_even_with_configured_key(self):
        summary = {'facts': [{'id': 'Fabcd', 'value': 18, 'unit': '篇'}], 'notes': []}
        factory = Mock(side_effect=AssertionError('Network call must not occur'))
        with patch.dict(os.environ, {'OPENAI_API_KEY': 'test-only-placeholder'}):
            self.assertEqual(load_config().api_key, '')
            record = ask_ai(summary, AIConfig(api_key='test-only-placeholder'), client_factory=factory)
        factory.assert_not_called()
        self.assertEqual(record['usage']['total_tokens'], 0)
        self.assertIn('规则演示', record['model'])

    def test_archive_roundtrip_and_invalid_archive(self):
        from cloud_projects import import_project, load_project
        from project_store import ProjectError
        app = AppTest.from_file(APP, default_timeout=90).run()
        app.button('cover-start').click().run()
        app.button('source-card-0').click().run()
        widget(app, 'button', '加载演示数据').click().run()
        widget(app, 'button', '保存项目').click().run()
        data = app.session_state['active_project']
        private = {}
        with patch('cloud_projects.projects', return_value=private):
            restored = import_project(json.dumps(data, ensure_ascii=False).encode('utf-8'))
            self.assertNotEqual(data['id'], restored['id'])
            self.assertEqual(restored['snapshot'], data['snapshot'])
            self.assertEqual(load_project(restored['id'])['state'], data['state'])
            with self.assertRaises(ProjectError):
                import_project(b'{}')
            with self.assertRaises(ProjectError):
                load_project(data['id'])


if __name__ == '__main__':
    unittest.main()
