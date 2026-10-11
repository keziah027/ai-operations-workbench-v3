"""导出当前已确认的数据和分析快照，在内存中生成，不保存登录资料。"""
from datetime import date, datetime
from io import BytesIO
import json
import math
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from content_analysis import safe_link
from note_schema import normalize_notes, COUNT_FIELDS
from performance import add_ratios

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
METRICS = ["点赞", "收藏", "评论"]
GREEN = "DCEEE1"
INK = "244633"


def export_filename(kind, name, captured_at):
    clean = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name or "账号")).strip(" .")[:40] or "账号"
    try:
        stamp = datetime.fromisoformat(captured_at).strftime("%Y%m%d_%H%M%S")
    except (ValueError, TypeError):
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{kind}_{clean}_数据与分析_{stamp}.xlsx"


def prepared_notes(notes):
    frame = pd.DataFrame(normalize_notes(notes)).copy()
    for field in ["标题", "发布时间", "笔记链接"]:
        if field not in frame:
            frame[field] = None
    for metric in COUNT_FIELDS:
        if metric not in frame:
            frame[metric] = None
        frame[metric] = pd.to_numeric(frame[metric], errors="coerce")
    if "序号" not in frame:
        frame.insert(0, "序号", range(1, len(frame) + 1))
    divisor = frame["点赞"].where(frame["点赞"].gt(0))
    frame["藏赞比"] = frame["收藏"] / divisor
    frame["评赞比"] = frame["评论"] / divisor
    frame["互动总量"] = frame[METRICS].sum(axis=1, min_count=3)
    frame = add_ratios(frame)
    # 只转换完整首发日期；不把“编辑于”等文字推断成首次发布时间。
    def source_date(value):
        if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}(?::\d{2})?)?", value):
            try:
                return datetime.fromisoformat(value)
            except ValueError:
                pass
        return value
    frame["发布时间"] = frame["发布时间"].map(source_date)
    return frame


def metric_summary(frame):
    return pd.DataFrame([{"指标": metric, "已确认合计": frame[metric].sum(min_count=1),
                          "平均值": frame[metric].mean(), "有效篇数": int(frame[metric].count()),
                          "本次笔记数": len(frame)} for metric in METRICS])


def ranking(frame):
    rows = []
    for metric in METRICS:
        top = frame.dropna(subset=[metric]).sort_values(metric, ascending=False, kind="stable").head(10)
        for position, (_, note) in enumerate(top.iterrows(), 1):
            rows.append({"榜单": f"{metric} TOP10", "排名": position, "序号": note["序号"],
                         "标题": note["标题"], "指标数值": note[metric], "发布时间": note["发布时间"],
                         "笔记链接": note["笔记链接"]})
    return pd.DataFrame(rows, columns=["榜单", "排名", "序号", "标题", "指标数值", "发布时间", "笔记链接"])


class Report:
    def __init__(self):
        self.book = Workbook()
        self.book.remove(self.book.active)
        self.book.properties.creator = "AI运营工作台"
        self.long_text = []

    def sheet(self, name):
        sheet = self.book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        return sheet

    def cell(self, sheet, row, column, value, header=""):
        if isinstance(value, (list, dict, tuple)):
            value = json.dumps(value, ensure_ascii=False, default=str)
        elif value is None or (not isinstance(value, str) and pd.isna(value)):
            value = None
        elif hasattr(value, "item"):
            value = value.item()
        if isinstance(value, float) and not math.isfinite(value):
            value = None
        if isinstance(value, pd.Timestamp):
            value = value.to_pydatetime()
        if isinstance(value, datetime) and value.tzinfo:
            value = value.isoformat()  # Excel 无时区类型，保留原始时区文本。
        if isinstance(value, str):
            value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
            if len(value) > 32000:
                for part, offset in enumerate(range(0, len(value), 32000), 1):
                    self.long_text.append({"原工作表": sheet.title, "原单元格": f"{get_column_letter(column)}{row}",
                                           "分段序号": part, "内容": value[offset:offset + 32000]})
                value = "完整内容见“长文本补充”，按原单元格和分段序号顺序阅读。"
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            cell.data_type = "s"  # 不把作者正文或人工输入中的 = 等内容当成 Excel 公式。
            if ("链接" in header or header == "来源") and safe_link(value):
                cell.hyperlink = value
                cell.font = Font(name="Microsoft YaHei", size=11, color="47795C", underline="single")
        else:
            cell.number_format = "#,##0.00" if isinstance(value, float) and not value.is_integer() else "#,##0"
        if isinstance(value, (datetime, date)):
            cell.number_format = "yyyy-mm-dd hh:mm" if isinstance(value, datetime) else "yyyy-mm-dd"
        if header in ["藏赞比", "评赞比", '观看曝光比', '平均藏赞比', '平均评赞比', '平均观看曝光比']:
            cell.number_format = "0.00%"
        cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="right" if isinstance(value, (int, float)) else "left")
        if not cell.hyperlink:
            cell.font = Font(name="Microsoft YaHei", size=11, color="263E30")
        return cell

    def table(self, sheet, title, frame, row=2, filter_rows=False):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(2, len(frame.columns)))
        title_cell = self.cell(sheet, row, 1, title)
        title_cell.font = Font(name="Microsoft YaHei", size=15, bold=True, color=INK)
        title_cell.alignment = Alignment(vertical="center", wrap_text=False)
        sheet.row_dimensions[row].height = 25
        header_row = row + 2
        frame = frame.copy()
        if len(frame.columns) == 0:
            frame = pd.DataFrame(columns=["说明"])
        percent_cols = set()
        for field in list(frame.columns):
            if "（%）" in str(field):
                renamed = str(field).replace("（%）", "")
                frame[field] = pd.to_numeric(frame[field], errors="coerce") / 100
                frame.rename(columns={field: renamed}, inplace=True)
                percent_cols.add(renamed)
        for col, name in enumerate(frame.columns, 1):
            cell = self.cell(sheet, header_row, col, str(name))
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color=INK)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="A8C7B3"))
            width = 18
            if any(word in str(name) for word in ["标题", "正文", "依据", "结论", "观察", "内容", "说明", "链接", "开头", "结尾"]):
                width = 46
            elif name in ["值", "记录"]:
                width = 80
            elif name == "发布时间":
                width = 24
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width or 0, width)
        sheet.row_dimensions[header_row].height = 34
        for offset, values in enumerate(frame.itertuples(index=False, name=None), 1):
            target = header_row + offset
            lines = 1
            for col, (name, value) in enumerate(zip(frame.columns, values), 1):
                cell = self.cell(sheet, target, col, value, str(name))
                if name in percent_cols:
                    cell.number_format = "0.00%"
                if isinstance(cell.value, str):
                    width = sheet.column_dimensions[get_column_letter(col)].width
                    logical_lines = cell.value.splitlines() or [""]
                    lines = max(lines, sum(max(1, math.ceil(len(line) / max(1, int(width / 1.7)))) for line in logical_lines))
            sheet.row_dimensions[target].height = min(180, max(24, 16 * lines + 8))
        last_row = header_row + len(frame)
        if frame.empty:
            self.cell(sheet, header_row + 1, 1, "本次没有可导出的记录；请查看分析汇总中的范围和停止原因。")
            sheet.row_dimensions[header_row + 1].height = 42
            last_row += 1
        if filter_rows:
            sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(frame.columns))}{max(header_row, header_row + len(frame))}"
            sheet.freeze_panes = f"C{header_row + 1}"
            sheet.print_title_rows = f"1:{header_row}"
        return last_row + 3, header_row

    def finish(self):
        if self.long_text:
            self.table(self.sheet("长文本补充"), "超长内容分段（按分段序号拼接）", pd.DataFrame(self.long_text), filter_rows=True)
        self.book.worksheets[0].sheet_properties.tabColor = "47795C"
        output = BytesIO()
        self.book.save(output)
        return output.getvalue()


def report_summary(report, snapshot, notes, kind, extra=()):
    sheet = report.sheet("分析汇总")
    metadata = [("数据来源", kind), ("读取时间", snapshot.get("captured_at")),
                ("导出时间", datetime.now().astimezone().isoformat(timespec="seconds")),
                ("本次读取篇数", len(notes)), ("设定上限", snapshot.get("requested_limit")),
                ("停止原因", snapshot.get("stop_reason", "未记录")),
                ("读取提示", "；".join(snapshot.get("warnings", [])) or "无"),
                *extra,
                ("导出方式", "当前数据与分析结果快照。修改 Excel 明细不会自动更新汇总；请在工作台重新分析并导出。"),
                ("空值与零", "空单元格为未读取或未确认，数字 0 为已确认的零。各项均值仅使用该项有效记录。"),
                ("互动率口径", "藏赞比=收藏/点赞，评赞比=评论/点赞；点赞为0或所需数值缺失时留空。不是曝光转化率。"),
                ("长文本", "正文等字段可在单元格编辑栏查看完整内容；超过 Excel 单格限制时放入长文本补充。")]
    following, _ = report.table(sheet, "AI运营工作台 · 数据与分析", pd.DataFrame(metadata, columns=["项目", "值"]))
    report.table(sheet, "本次笔记互动汇总", metric_summary(notes), row=following)


def public_workbook(snapshot, analysis):
    notes = prepared_notes(snapshot.get("notes", []))
    report = Report()
    report_summary(report, snapshot, notes, snapshot.get('source_kind', '小红书公开账号'), [
        ("账号", snapshot.get("name")), ("来源", snapshot.get("url")), ("账号简介", snapshot.get("bio", "")),
        ("主页指标原文", snapshot.get("profile_text", "")), ("详情成功篇数", snapshot.get("successful_details", 0)),
        ("样本范围", snapshot.get('scope', '仅本次读取到的笔记，不代表账号全部历史。')),
        ("支柱互动率口径", "支柱藏赞比、评赞比为逐篇有效比值的平均值；点赞为0或数据缺失时不参与，各列提供有效篇数。")])
    v2 = 'profile' in analysis
    five = report.sheet("五模块分析结果" if v2 else "五维分析结果")
    profile_rows = []
    if v2:
        profile_rows = [("身份初判", analysis['profile']['identity']), ("身份初判置信度", analysis['profile']['identity_confidence']),
                        ("领域初判", '、'.join(analysis['profile']['domains'])),
                        ("置信度含义", analysis['profile']['method']), ("用户确认身份", analysis['identity']),
                        ("用户确认领域", '、'.join(analysis['domains'])), ("用户确认目标", '、'.join(analysis['goals']) or '尚未确认'),
                        ("商品分析启用依据", analysis['commerce_reason']), ('本次分析重点','、'.join(analysis.get('analysis_goals',[]))),
                        ("类型识别依据", analysis['profile']['reason']), ("人工诊断补充", analysis['diagnosis_note'])]
    row, _ = report.table(five, "账号画像", pd.DataFrame([
        *profile_rows,
        ("账号定位结论", analysis["position"]), ("核心人群假设及待验证问题", analysis["audience"]),
        ("人群口径", "兴趣和需求假设，不是已确认的粉丝年龄、性别、购买力画像。"),
        ("高互动比较指标", analysis["metric"]), ("该指标有效篇数", len(analysis["valid"])),
        ("高互动组篇数", len(analysis["top"])),
        ("高互动口径", "有效样本前25%，向上取整，同值按读取顺序。至少8篇且数值有差异才比较特征。"),
        ("封面口径", "视觉类型由人工标注；对比只使用已标注样本，两组各至少2篇。"),
        (("商业化口径", "仅公开文字线索与已识别链接，不能确认点击、加购、订单或GMV。")
         if not v2 or analysis['commerce_enabled'] else ("诊断口径", "按身份、领域、用户确认目标与可用数据选择重点，不推断真实画像、信任或因果。"))], columns=["项目", "记录"]))
    sections = [("人群假设与依据", analysis["audience_findings"]), ("内容支柱占比", analysis["summary"]),
                ("高互动规律", analysis["comparisons"])]
    if v2:
        sections += [('最终主题与代表笔记', analysis.get('final_topics',pd.DataFrame())), ("身份与领域候选依据", analysis['profile']['candidates']),
                     ("主题自动归纳依据（人工修订后以支柱统计为准）", analysis['topics']), ("运营诊断", analysis['diagnosis'])]
    else:
        sections += [("商业化承接线索", pd.DataFrame(analysis["commerce"]))]
    for title, frame in sections:
        row, header = report.table(five, title, frame, row=row)
        if title == "内容支柱占比" and not frame.empty:
            chart = BarChart()
            chart.title = "本次样本的内容支柱"
            chart.y_axis.title = "篇数"
            chart.add_data(Reference(five, min_col=2, min_row=header, max_row=header + len(frame)), titles_from_data=True)
            chart.set_categories(Reference(five, min_col=1, min_row=header + 1, max_row=header + len(frame)))
            chart.width, chart.height = 19, 9
            five.add_chart(chart, f"P{header}")
    if analysis.get('packaging_tables'):
        for title,table in analysis['packaging_tables'].items():
            row,_=report.table(five,title,table,row=row)
        row,_=report.table(five,'封面类型配置',pd.DataFrame({'主类型':analysis['editing']['packaging']['cover_types']}),row=row)
    top_columns = [column for column in ["序号", "标题", "主支柱", analysis["metric"], "笔记链接"] if column in analysis["top"]]
    row, _ = report.table(five, "高互动组笔记", analysis["top"][top_columns], row=row)
    annotated = analysis["detail"]
    covers = annotated[annotated["封面类型"] != "待标注"]["封面类型"].value_counts().rename_axis("封面类型").reset_index(name="篇数")
    report.table(five, "已标注封面分布", covers, row=row)
    detail = annotated.drop(columns=["正文", "点赞", "收藏", "评论", "互动总量", "可见入口"], errors="ignore").copy()
    detail["分类依据"] = analysis["annotations"]["分类依据"].to_list()
    if v2 and not analysis['commerce_enabled']:
        detail = detail.drop(columns=['商品提及', '购买引导', '店铺/橱窗', '商品/系列（人工确认）'], errors='ignore')
    report.table(report.sheet("内容拆解明细"), "内容结构、包装特征、证据与人工标注", detail, filter_rows=True)
    report.table(report.sheet("排行榜"), "本次笔记三项 TOP10", ranking(notes), filter_rows=True)
    report.table(report.sheet("笔记明细"), "已读取笔记及原始显示字段", notes, filter_rows=True)
    if v2:
        report.table(report.sheet('比较样本范围'), '每篇笔记的纳入与排除原因', analysis['sample_audit'], filter_rows=True)
        for table_name, table in analysis['backend_tables'].items():
            if table_name=='数据完整度' or annotated[['曝光','观看','涨粉']].notna().any().any():
                report.table(report.sheet(table_name), table_name, table, filter_rows=True)
    if 'audience_all' in analysis:
        report.table(report.sheet('人群假设复核记录'),'自动证据与人工处理记录（最终采用见五模块）',analysis['audience_all'],filter_rows=True)
    if analysis.get('ai'):
        from ai_report import report_tables
        for title,table in report_tables(analysis['ai']).items():
            report.table(report.sheet(title),title,table,filter_rows=True)
    if snapshot.get('creator_data'):
        creator = snapshot['creator_data']
        overview = creator.get('overview') or {}
        report.table(report.sheet('账号期间总览'), overview.get('period', '账号期间总览'), pd.DataFrame(overview.get('metrics', [])), filter_rows=True)
    return report.finish()


def creator_workbook(snapshot):
    notes = prepared_notes(snapshot.get("notes", []))
    report = Report()
    overview = snapshot.get("overview") or {}
    filters = snapshot.get("filters") or {}
    report_summary(report, snapshot, notes, "小红书创作服务平台", [
        ("账号", overview.get("name")), ("小红书账号", str(overview.get("account_id", ""))),
        ("来源", snapshot.get("source")), ("账号总览周期", overview.get("period")),
        ("笔记首发开始日期", filters.get("start")), ("笔记首发结束日期", filters.get("end")),
        ("笔记类型筛选", filters.get("type")), ("后台符合筛选篇数", snapshot.get("platform_total")),
        ("已读完筛选范围", "是" if snapshot.get("complete_filtered") else "否"),
        ("统计范围", "账号期间总览与本次笔记汇总口径不同，分表展示。")])
    report.table(report.sheet("账号期间总览"), overview.get("period", "账号期间总览"), pd.DataFrame(overview.get("metrics", [])), filter_rows=True)
    report.table(report.sheet("排行榜"), "本次笔记三项 TOP10", ranking(notes), filter_rows=True)
    trend_source = notes.dropna(subset=["发布时间", "点赞"]).copy()
    trend_source["发布时间"] = pd.to_datetime(trend_source["发布时间"], errors="coerce")
    trend_source = trend_source.dropna(subset=["发布时间"])
    trend_source["发布日期"] = trend_source["发布时间"].dt.date
    trend = trend_source.groupby("发布日期", as_index=False)["点赞"].mean().rename(columns={"点赞": "平均点赞"})
    trend_sheet = report.sheet("发布时间趋势")
    _, header = report.table(trend_sheet, "按首发日期比较当前平均点赞（不是每日新增点赞）", trend, filter_rows=True)
    if not trend.empty:
        chart = LineChart()
        chart.title = "按首发日期比较平均点赞"
        chart.y_axis.title = "平均点赞"
        chart.x_axis = DateAxis(crossAx=100)
        chart.x_axis.number_format = "mm-dd"
        chart.x_axis.title = "首发日期"
        chart.add_data(Reference(trend_sheet, min_col=2, min_row=header, max_row=header + len(trend)), titles_from_data=True)
        chart.set_categories(Reference(trend_sheet, min_col=1, min_row=header + 1, max_row=header + len(trend)))
        chart.width, chart.height = 21, 10
        trend_sheet.add_chart(chart, "D4")
    report.table(report.sheet("笔记明细"), "后台筛选范围内已读取的笔记", notes, filter_rows=True)
    return report.finish()
